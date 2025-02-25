"""Utilities for the rankings module"""
from datetime import datetime
import openpyxl
from rankings import logging_config

logger = logging_config.logger
LOAD_WAIT_SECONDS = 10

def get_core_url_part(url: str) -> str:
    """Extracts the identifying part of a URL. E.g. In:
    http://r.statista.com/best-employers-singapore-2022/

    The identifying part is 'best-employers-singapore-2022'

    Arguments:
        url (str) -- any url that we suspect contains ranking data

    Returns:
        The part of the URL that identifies the content

    """
    parts = [part for part in url.split('/') if part != '']
    # print(parts)

    # If the last part is 'ranking', then the second to last part is the identifier
    if parts[-1] == 'ranking':
        core_part = parts[-2]
    else:
        # It's a non-standard URL, look for the part based on keywords
        parts = (part for part in parts if 'best' in part or 'employers-' in part)
        try:
            core_part = next(parts)
        except StopIteration as exc:
            error_message = f'Error: No matching URL part found in {url}'
            raise ValueError(error_message) from exc
        try:
            next(parts)
            raise ValueError(
                f'Error: More than one matching URL part found in {url}')
        except StopIteration:
            pass

    return core_part

def predict_country_study_year(core_part, country_map, study_map):
    """Tokenizes the core part of a url and attempts to predict the country, study and year.
    
    Arguments:
        core_part (str) -- the part of the url that comes after the domain name and before 
            the ranking. Provided by get_core_url_part.
        country_map (dict) -- maps potential country tokens to country names
            retrieved from datasets.xlsx sheet 'country_map'
        study_map (dict) -- maps potential study tokens to study names
            retrieved from datasets.xlsx sheet 'study_map'

    Returns:
        tuple -- (country, study, year) if all three can be predicted, else
        None
    
    """

    logger.info('Attempting to predict variables for core_part: %s', core_part)

    tokens = core_part.split('-')
    logger.info('Tokens: %s', tokens)

    # TODO: implement translation of tokens to English
    exclusions = ["best", "beste", "employers", "arbeitgeber", "the", "feur"]

    # remove exclusions from tokens
    tokens = [token for token in tokens if token not in exclusions]

    logger.info('Tokens without exclusions: %s', tokens)


    # if one of the tokens is a four-digit integer, it's probably the year
    for token in tokens:
        if len(token) == 4 and token.isdigit():
            predicted_year = int(token)
            tokens.remove(token)
            break
    else:
        logger.warning('Unable to predict year for %s', core_part)
        return None

    logger.info('Predicted year: %s', predicted_year)

    # if one of the tokens is a country name that also appears in the table, it's probably
    # the country
    for token in tokens:
        if token in country_map:
            predicted_country = country_map[token]
            tokens.remove(token)
            # else if the token ends in an s, remove it and try again
        elif token.endswith('s'):
            token = token[:-1]
            if token in country_map:
                predicted_country = country_map[token]
                break
        else:
            logger.warning('Unable to predict country for %s', core_part)
            return None

    logger.info('Predicted country: %s', predicted_country)

    # if one of the tokens is a study name that also appears in the table, it's probably the study
    if len(tokens) == 0:
        predicted_study = 'best'
    else:
        for token in tokens:
            if token in study_map:
                predicted_study = study_map[token]
                tokens.remove(token)
                break
        else:
            logger.warning('Unable to predict study for %s', core_part)
            return None

    logger.info('Predicted study: %s', predicted_study)

    # if we have a year, country and study, we're done
    if predicted_year and predicted_country and predicted_study:
        return (predicted_country, predicted_study, predicted_year)
    else:
        logger.warning('Unable to predict country, study and year for %s', core_part)
        return None

def add_new_dataset(data_dict: dict):
    """Adds a new row to the datasets sheet in datasets.xlsx
    
    Arguments:
        data_dict (dict) -- the country, study, year and url as keys of a dictionary
            the dict values are the values to be written to the row

    Returns:
        list representing the new row. Includes datetime the row was added
    
    
    """

    filepath = r'..\data\datasets.xlsx'

    # Load the workbook and select the 'datasets' sheet
    workbook = openpyxl.load_workbook(filepath)
    sheet = workbook['datasets']
    table = sheet.tables['datasets']

    data_list = []
    for col in table.tableColumns:
        formula = col.calculatedColumnFormula
        if col.name in data_dict:
            data_list.append(data_dict[col.name])
        elif col.name == 'added':
            data_list.append(datetime.now())
        elif col.name == 'link_valid':
            data_list.append(1)
        elif not formula is None:
            data_list.append('=' + formula.text)
        else:
            data_list.append(None)

    sheet.append(data_list)
    logger.info('New row added to dataset repository: %s', data_list)

    # expand table reference
    tl, br = table.ref.split(':')
    alpha = ''.join(c for c in br if c.isalpha())
    table.ref = f'{tl}:{alpha}{sheet.max_row}'

    # Save the workbook
    workbook.save(filepath)

    return data_list
